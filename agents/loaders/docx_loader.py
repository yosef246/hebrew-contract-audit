# docx_loader.py — חילוץ טקסט טהור מ-.docx (python-docx).
# post-F: יחובר ל-pipeline האמיתי לצד ה-PDF extractor. כרגע פונקציה עצמאית אחת.
from docx import Document

def extract_docx(path: str) -> str:
    """טקסט טהור מ-.docx: פסקאות + תאי טבלה (חוזים רבים משתמשים בטבלאות), מחוברים בשורות."""
    doc = Document(path)
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts).strip()
